"""把生成的 Markdown 文档转成 .docx（python-docx）。

两个入口：
- markdown_to_docx(md, title)  —— 旧的简单转换（保留兼容）。
- render_gwdocx(md, meta)      —— 公文格式渲染：封面/页眉页脚/宋体小四正文/标题居中/
  章节公文样式/表格底纹/mermaid 插图。网页 docx-preview 预览与下载同用此产物。
"""
import re
from io import BytesIO

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.core.mermaid_render import render_mermaid_to_png


CJK_FONT = "宋体"          # Word 里中文显示；Linux 上无宋体时按字体回退仍可读
HEADING_FONT = "黑体"


def _set_cell_text(cell, text: str):
    cell.text = ""
    para = cell.paragraphs[0]
    _add_inline(para, text.strip())


def _apply_cjk(run, font_name: str):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def _add_inline(para, text: str, font_name: str = CJK_FONT, size: int = None, bold: bool = False):
    """处理 **加粗** 行内标记，逐段加 run。"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = bold
        content = part
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            is_bold = True
            content = part[2:-2]
        run = para.add_run(content)
        run.bold = is_bold
        if size:
            run.font.size = Pt(size)
        _apply_cjk(run, font_name)


def _is_table_sep(line: str) -> bool:
    # |---|---| 这种分隔行
    s = line.strip()
    return bool(re.match(r"^\|?\s*:?-{2,}.*$", s)) and set(s) <= set("|-: \t")


def _parse_table_row(line: str):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _add_scaled_picture(paragraph, png_bytes, max_w_in=5.0, max_h_in=6.0):
    """按图片宽高比等比缩放插入：竖图受高度上限约束、横图受宽度上限约束，
    避免竖向流程图等比放大后过高撑满/跨页。"""
    import struct
    # 读 PNG 尺寸(IHDR: 宽高在 offset 16-24)
    w_px = h_px = 0
    try:
        if png_bytes[:8] == b"\x89PNG\r\n\x1a\n":
            w_px, h_px = struct.unpack(">II", png_bytes[16:24])
    except Exception:
        pass
    run = paragraph.add_run()
    if not w_px or not h_px:
        run.add_picture(BytesIO(png_bytes), width=Inches(min(max_w_in, 5.0)))
        return
    # 图片默认 96dpi 换算英寸
    w_in, h_in = w_px / 96.0, h_px / 96.0
    scale = min(max_w_in / w_in, max_h_in / h_in, 1.0)   # 不放大,只缩小
    run.add_picture(BytesIO(png_bytes), width=Inches(w_in * scale))


def _emit_code_block(doc, code: str):
    """代码块降级输出：等宽字体逐行（mermaid 渲染失败或普通代码块用）。"""
    for ln in code.split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(ln)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def markdown_to_docx(md: str, title: str = "") -> bytes:
    doc = Document()

    # 默认正文字体
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = CJK_FONT

    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 跳过结束标记残留
        if "<!--DOC_END-->" in stripped:
            stripped = stripped.replace("<!--DOC_END-->", "").strip()
            if not stripped:
                i += 1
                continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 水平线
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            i += 1
            continue

        # 代码块 ```lang ... ```
        m = re.match(r"^```+\s*([A-Za-z0-9_-]*)\s*$", stripped)
        if m:
            lang = (m.group(1) or "").lower()
            i += 1
            code_lines = []
            while i < n and not re.match(r"^```+\s*$", lines[i].strip()):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过收尾 ```
            code = "\n".join(code_lines)
            if lang == "mermaid":
                png = render_mermaid_to_png(code)
                if png:
                    from io import BytesIO as _BIO
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    try:
                        run.add_picture(_BIO(png), width=Inches(5.5))
                    except Exception:
                        _emit_code_block(doc, code)
                else:
                    _emit_code_block(doc, code)   # 渲染失败：降级为源码文本
            else:
                _emit_code_block(doc, code)
            continue

        # 标题 # ~ ######
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_paragraph()
            sizes = {1: 18, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}
            _add_inline(p, text, font_name=HEADING_FONT, size=sizes.get(level, 11), bold=True)
            i += 1
            continue

        # 表格：当前行是 |...| 且下一行是分隔行
        if stripped.startswith("|") and i + 1 < n and _is_table_sep(lines[i + 1]):
            header = _parse_table_row(lines[i])
            i += 2  # 跳过表头和分隔行
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_parse_table_row(lines[i]))
                i += 1
            ncol = len(header)
            table = doc.add_table(rows=1, cols=ncol)
            table.style = "Table Grid"
            for j, h in enumerate(header):
                if j < ncol:
                    _set_cell_text(table.rows[0].cells[j], h)
            for r in rows:
                cells = table.add_row().cells
                for j in range(ncol):
                    _set_cell_text(cells[j], r[j] if j < len(r) else "")
            continue

        # 引用块 >
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_inline(p, text, size=10)
            i += 1
            continue

        # 有序列表 1. / 1、
        m = re.match(r"^(\d+)[\.、]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, m.group(2).strip())
            i += 1
            continue

        # 无序列表 - / * / •
        m = re.match(r"^[-*•]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(1).strip())
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _add_inline(p, stripped)
        i += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════
# 公文格式渲染 render_gwdocx —— 网页预览(docx-preview)与下载同用此产物
# ══════════════════════════════════════════════════════════════════════════

# 公文标题层级样式：markdown #层级 → (字号pt, 是否居中, 字体)
_GW_HEADING = {
    1: (16, True, HEADING_FONT),    # 一级：黑体三号居中（章）
    2: (14, False, HEADING_FONT),   # 二级：黑体四号（节）
    3: (12, False, HEADING_FONT),   # 三级：黑体小四
    4: (12, False, CJK_FONT),       # 四级及以下：宋体小四加粗
    5: (12, False, CJK_FONT),
    6: (12, False, CJK_FONT),
}


def _shade_cell(cell, fill="D9E2F3"):
    """给单元格加底纹（表头用淡蓝）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _hdr_cell(cell, text, bold=False, size=9, center=True):
    cell.text = ""
    p = cell.paragraphs[0]
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text or "")
    r.bold = bold; r.font.size = Pt(size); _apply_cjk(r, CJK_FONT)


def _table_borders(table, sz="6", color="000000"):
    """给表格设细黑边框(白底黑线,替代 Table Grid 的默认深色渲染)。"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    # 清掉可能的 tblStyle
    for st in tblPr.findall(qn("w:tblStyle")):
        tblPr.remove(st)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _set_header_footer(doc, meta):
    """不插页眉表头（用户要求，觉得不好看）。仅在页脚居中放页码。"""
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True   # 封面页(首页)页脚也不显示页码
    sec.first_page_footer.paragraphs[0].text = ""
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(); r.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)



def _cover(doc, meta):
    """封面/抬头：标题居中大字 + 编号/版本/日期/受控 信息表。"""
    title = meta.get("title", "文件")
    company = meta.get("company", "")
    # 顶部留白
    for _ in range(2):
        doc.add_paragraph()
    # 公司名（抬头，小二居中）
    if company:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(company); r.bold = True; r.font.size = Pt(18); _apply_cjk(r, HEADING_FONT)
    # 中部留白
    for _ in range(3):
        doc.add_paragraph()
    # 文件标题（黑体特大居中）
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(30); _apply_cjk(r, HEADING_FONT)
    # 标题到信息表的留白
    for _ in range(3):
        doc.add_paragraph()
    # 信息表：编号/版本/生效日期/密级/受控（居中、收窄、浅灰表头）
    info = [
        ("文件编号", meta.get("doc_no", "")),
        ("版本号", meta.get("version", "")),
        ("生效日期", meta.get("date", "")),
        ("密　　级", meta.get("secrecy", "内部")),
        ("受控状态", meta.get("controlled", "受控")),
    ]
    tbl = doc.add_table(rows=len(info), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_borders(tbl)
    for idx, (k, v) in enumerate(info):
        kc, vc = tbl.rows[idx].cells[0], tbl.rows[idx].cells[1]
        kc.width = Inches(1.6); vc.width = Inches(3.0)
        _set_cell_text(kc, k)
        for run in kc.paragraphs[0].runs: run.bold = True
        kc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_text(vc, v or "　")
        _shade_cell(kc, "F2F2F2")
    # 分页，正文从新页开始
    doc.add_page_break()


def render_gwdocx(md: str, meta: dict = None) -> bytes:
    """公文格式渲染：Markdown → 规范 docx。meta: {title, company, doc_no, version, date, controlled}"""
    meta = meta or {}
    doc = Document()

    # 页边距（公文常用上下2.54 左右3.17→用默认略调）
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.1); sec.right_margin = Inches(1.1)

    # 正文默认：宋体小四(12pt)，1.5倍行距
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12); normal.font.name = CJK_FONT
    normal.paragraph_format.line_spacing = 1.5

    # 封面 + 页眉页脚
    _cover(doc, meta)
    _set_header_footer(doc, meta)

    lines = md.replace("\r\n", "\n").split("\n")
    i, n = 0, len(lines)
    while i < n:
        stripped = lines[i].strip()
        if "<!--DOC_END-->" in stripped:
            stripped = stripped.replace("<!--DOC_END-->", "").strip()
            if not stripped:
                i += 1; continue
        if not stripped:
            i += 1; continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            i += 1; continue

        # 代码块 / mermaid
        m = re.match(r"^```+\s*([A-Za-z0-9_-]*)\s*$", stripped)
        if m:
            lang = (m.group(1) or "").lower(); i += 1
            code_lines = []
            while i < n and not re.match(r"^```+\s*$", lines[i].strip()):
                code_lines.append(lines[i]); i += 1
            i += 1
            code = "\n".join(code_lines)
            if lang == "mermaid":
                png = render_mermaid_to_png(code)
                if png:
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    try:
                        _add_scaled_picture(p, png, max_w_in=5.0, max_h_in=6.0)
                    except Exception: _emit_code_block(doc, code)
                else: _emit_code_block(doc, code)
            else: _emit_code_block(doc, code)
            continue

        # 标题（公文样式：居中/黑体/字号按层级）
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1)); text = m.group(2).strip()
            size, center, font = _GW_HEADING.get(level, (12, False, CJK_FONT))
            p = doc.add_paragraph()
            if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(4)
            _add_inline(p, text, font_name=font, size=size, bold=True)
            i += 1; continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and _is_table_sep(lines[i + 1]):
            header = _parse_table_row(lines[i]); i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_parse_table_row(lines[i])); i += 1
            ncol = len(header)
            table = doc.add_table(rows=1, cols=ncol)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            _table_borders(table)
            # 按内容长度自适应列宽：统计每列最大文字长度→按比例分配总宽(6.0英寸)
            col_len = [max(len(header[j]) if j < len(header) else 0,
                           *[len(r[j]) if j < len(r) else 0 for r in rows] if rows else [0])
                       for j in range(ncol)]
            col_len = [max(2, c) for c in col_len]        # 每列至少权重2
            total = sum(col_len) or ncol
            TOTAL_IN = 6.0
            widths = [max(0.8, min(4.2, TOTAL_IN * c / total)) for c in col_len]  # 每列0.8~4.2英寸
            def _setw(cell, w):
                cell.width = Inches(w)
            for j, h in enumerate(header):
                if j < ncol:
                    _set_cell_text(table.rows[0].cells[j], h)
                    _shade_cell(table.rows[0].cells[j])
                    for run in table.rows[0].cells[j].paragraphs[0].runs: run.bold = True
                    _setw(table.rows[0].cells[j], widths[j])
            for rrow in rows:
                cells = table.add_row().cells
                for j in range(ncol):
                    _set_cell_text(cells[j], rrow[j] if j < len(rrow) else "")
                    _setw(cells[j], widths[j])
            doc.add_paragraph()
            continue

        # 引用块
        if stripped.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(21)
            _add_inline(p, stripped.lstrip(">").strip(), size=11)
            i += 1; continue

        # 有序/无序列表
        m = re.match(r"^(\d+)[\.、]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number"); _add_inline(p, m.group(2).strip(), size=12)
            i += 1; continue
        m = re.match(r"^[-*•]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet"); _add_inline(p, m.group(1).strip(), size=12)
            i += 1; continue

        # 正文段落：首行缩进2字符
        p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(24)
        _add_inline(p, stripped, size=12)
        i += 1

    buf = BytesIO(); doc.save(buf)
    return buf.getvalue()
